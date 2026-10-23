"""
Leitura da planilha "Distribuição de Funções (Pastas): Encontro e Pós-Encontro".

Layout esperado (o mesmo da planilha feita nas reuniões):
    Função | Encontro | Pós-Encontro | Coordenador | Membro de Equipe | Observações | Coordenador | Membro de Equipe | Observações
                                       └──────────── ENCONTRO ──────────────────┘   └────────── PÓS-ENCONTRO ────────────┘
Cada pasta ocupa uma ou mais linhas (células mescladas). Uma célula pode ter
vários nomes, um por linha.

Fica de fora (não é servo do sistema): células vermelhas (Apostolado), amarelas
ou laranjas (Casal da Missionária Jovem), "Xxx e Xxx" e textos com casal/casais.
Aqui só se lê a planilha; quem confere pastas e servos é o router.
"""
import io
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Optional

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

MOMENTOS = ("encontro", "pos")

# Cores de quem não é servo do sistema (RGB, sem o alfa)
_CORES_IGNORADAS = {"FF0000", "C00000", "FFC000", "FFFF00", "FFD966", "F4B084", "FFC7CE"}
_TEXTO_IGNORADO = re.compile(r"^x+ e x+$|\bcasa(l|is)\b|\bapostolado\b|\bmission[aá]ri", re.IGNORECASE)


class PlanilhaInvalida(ValueError):
    """A planilha não tem o formato esperado."""


@dataclass
class EquipeLida:
    coordenadores: list[str] = field(default_factory=list)
    membros: list[str] = field(default_factory=list)
    ignorados: list[str] = field(default_factory=list)
    marcada: bool = False  # tem o "X" na coluna Encontro/Pós-Encontro


@dataclass
class PastaLida:
    funcao: str   # como está escrito na planilha
    linha: int    # linha da planilha onde a pasta começa
    equipes: dict[str, EquipeLida] = field(default_factory=lambda: {m: EquipeLida() for m in MOMENTOS})


def normalizar(texto: str) -> str:
    """Sem acento, minúsculo, sem pontuação e com um espaço entre as palavras."""
    sem_acento = unicodedata.normalize("NFKD", texto)
    sem_acento = "".join(c for c in sem_acento if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", sem_acento.lower()).split())


def _texto(valor) -> str:
    return str(valor).strip() if valor is not None else ""


def _cor(cell: Cell) -> Optional[str]:
    try:
        fill = cell.fill
        if fill is None or fill.fill_type in (None, "none"):
            return None
        rgb = fill.fgColor.rgb if fill.fgColor is not None else None
        return rgb[-6:].upper() if isinstance(rgb, str) and len(rgb) >= 6 else None
    except (AttributeError, TypeError):
        return None


def _ignorado(cell: Cell, nome: str) -> bool:
    return _cor(cell) in _CORES_IGNORADAS or bool(_TEXTO_IGNORADO.search(nome))


def _nomes(valor) -> list[str]:
    """Uma célula pode trazer vários nomes, um por linha (ou separados por ;)."""
    partes = re.split(r"[\n;]+", _texto(valor))
    return [" ".join(p.split()) for p in partes if p.strip()]


def _achar_cabecalho(ws: Worksheet) -> tuple[int, int, dict[str, int], dict[str, dict[str, int]]]:
    """
    Linha do cabeçalho, coluna da Função, colunas do X (encontro/pos) e as
    colunas Coordenador / Membro / Observações de cada momento.
    """
    for linha in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        textos = {c.column: normalizar(_texto(c.value)) for c in linha if c.value is not None}
        col_funcao = next((col for col, t in textos.items() if t == "funcao"), None)
        if col_funcao is None:
            continue

        coords = sorted(col for col, t in textos.items() if t == "coordenador")
        membros = sorted(col for col, t in textos.items() if t.startswith("membro"))
        obs = sorted(col for col, t in textos.items() if t.startswith("observa"))
        if len(coords) < 2 or len(membros) < 2:
            raise PlanilhaInvalida(
                "Não achei as colunas Coordenador e Membro de Equipe do Encontro e do Pós-Encontro.")

        flags = {}
        for col, t in textos.items():
            if col >= coords[0]:
                continue
            if t in ("pos encontro", "posencontro", "pos"):
                flags["pos"] = col
            elif t == "encontro":
                flags["encontro"] = col

        colunas = {}
        for i, momento in enumerate(MOMENTOS):
            fim = coords[i + 1] if i + 1 < len(coords) else ws.max_column + 1
            colunas[momento] = {
                "coordenador": coords[i],
                "membro": membros[i],
                "obs": next((c for c in obs if coords[i] < c < fim), None),
            }
        return linha[0].row, col_funcao, flags, colunas

    raise PlanilhaInvalida('Não achei a linha de cabeçalho com a coluna "Função".')


def ler_planilha(conteudo: bytes) -> list[PastaLida]:
    try:
        wb = load_workbook(io.BytesIO(conteudo), data_only=True)
    except Exception as e:
        raise PlanilhaInvalida("Não consegui abrir o arquivo. Envie a planilha em .xlsx.") from e

    ws = wb.active
    linha_cab, col_funcao, flags, colunas = _achar_cabecalho(ws)
    primeira_col_equipe = min(c["coordenador"] for c in colunas.values())

    # Célula mesclada → intervalo (só a primeira célula guarda o valor)
    mescladas = {}
    for intervalo in ws.merged_cells.ranges:
        mescladas[(intervalo.min_row, intervalo.min_col)] = intervalo

    def funcao_da_linha(r: int) -> str:
        # O número da pasta pode estar numa coluna própria, antes do nome
        partes = []
        for col in range(1, (flags and min(flags.values())) or primeira_col_equipe):
            t = _texto(ws.cell(row=r, column=col).value)
            if t:
                partes.append(t)
        texto = " ".join(partes)
        return re.sub(r"^\d+\s*[-.)]?\s*", "", texto).strip()

    pastas: list[PastaLida] = []
    atual: Optional[PastaLida] = None
    obs_atual: dict[str, str] = {}

    for r in range(linha_cab + 1, ws.max_row + 1):
        funcao = funcao_da_linha(r)
        if funcao and normalizar(funcao) not in ("encontro", "pos encontro"):
            atual = PastaLida(funcao=funcao, linha=r)
            pastas.append(atual)
            obs_atual = {}
            for momento, col in flags.items():
                if normalizar(_texto(ws.cell(row=r, column=col).value)) == "x":
                    atual.equipes[momento].marcada = True
        if atual is None:
            continue

        for momento, cols in colunas.items():
            equipe = atual.equipes[momento]
            if cols["obs"]:
                obs = _texto(ws.cell(row=r, column=cols["obs"]).value)
                if obs:
                    obs_atual[momento] = obs

            for papel in ("coordenador", "membro"):
                cell = ws.cell(row=r, column=cols[papel])
                if cell.value is None:
                    continue
                destino = equipe.coordenadores if papel == "coordenador" else equipe.membros

                # Nome centralizado sobre Coordenador + Membro (pasta sem equipe):
                # é coordenador só se a observação disser só "coordenador(es)"
                intervalo = mescladas.get((r, cell.column))
                if papel == "coordenador" and intervalo and intervalo.max_col >= cols["membro"]:
                    obs = normalizar(obs_atual.get(momento, "") or _obs_abaixo(ws, r, cols["obs"]))
                    so_coordenadores = "coordenador" in obs and "+" not in obs_atual.get(momento, "")
                    destino = equipe.coordenadores if so_coordenadores else equipe.membros

                for nome in _nomes(cell.value):
                    if _ignorado(cell, nome):
                        equipe.ignorados.append(nome)
                    elif nome not in destino:
                        destino.append(nome)

    return pastas


def _obs_abaixo(ws: Worksheet, r: int, col_obs: Optional[int]) -> str:
    """A observação às vezes fica na linha de baixo da pasta (célula mesclada na vertical)."""
    if not col_obs:
        return ""
    for intervalo in ws.merged_cells.ranges:
        if intervalo.min_col <= col_obs <= intervalo.max_col and intervalo.min_row <= r <= intervalo.max_row:
            return _texto(ws.cell(row=intervalo.min_row, column=intervalo.min_col).value)
    return ""
