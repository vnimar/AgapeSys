"""Importação da distribuição de pastas (.xlsx). Planilhas fictícias, no mesmo formato da real."""
import io
from datetime import date

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from resources.bd.db import get_connection
from resources.importacao.pastas_xlsx import ler_planilha

ANO = date.today().year + 1
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VERMELHO = PatternFill("solid", fgColor="FFFF0000")
AMARELO = PatternFill("solid", fgColor="FFFFC000")


def planilha(pastas: list[dict]) -> bytes:
    """
    pastas: {"funcao", "e"/"p" (X), "ec"/"em"/"eo" (encontro: coord, membros, obs),
             "pc"/"pm"/"po" (pós), "mesclar": {"e"|"p"}, "vermelho"/"amarelo": [nomes]}
    Coordenador e membros: listas; cada nome numa linha da planilha.
    """
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "DISTRIBUIÇÃO DE FUNÇÕES (PASTAS): ENCONTRO E PÓS-ENCONTRO"
    ws["E2"], ws["H2"] = "ENCONTRO", "PÓS ENCONTRO"
    for col, texto in zip("ACDEFGHIJ", ["Função", "Encontro", "Pós-Encontro", "Coordenador", "Membro de Equipe",
                                        "Observações", "Coordenador", "Membro de Equipe", "Observações"]):
        ws[f"{col}3"] = texto
    ws.merge_cells("A3:B3")

    r = 4
    for numero, p in enumerate(pastas, start=1):
        altura = max(1, *(len(p.get(k, [])) for k in ("ec", "em", "pc", "pm")))
        ws[f"A{r}"], ws[f"B{r}"] = numero, p["funcao"]
        ws[f"C{r}"] = "X" if p.get("e") else None
        ws[f"D{r}"] = "X" if p.get("p") else None
        for prefixo, (cc, cm, co) in (("e", "EFG"), ("p", "HIJ")):
            ws[f"{co}{r}"] = p.get(prefixo + "o")
            if prefixo in p.get("mesclar", set()):
                # Nomes centralizados sobre Coordenador + Membro, numa célula só
                nomes = p.get(prefixo + "c", []) + p.get(prefixo + "m", [])
                ws[f"{cc}{r}"] = "\n".join(nomes) or None
                ws.merge_cells(f"{cc}{r}:{cm}{r}")
                continue
            for i, nome in enumerate(p.get(prefixo + "c", [])):
                ws[f"{cc}{r + i}"] = nome
            for i, nome in enumerate(p.get(prefixo + "m", [])):
                cell = ws[f"{cm}{r + i}"]
                cell.value = nome
                if nome in p.get("vermelho", []):
                    cell.fill = VERMELHO
                if nome in p.get("amarelo", []):
                    cell.fill = AMARELO
        if altura > 1:
            ws.merge_cells(f"B{r}:B{r + altura - 1}")
        r += altura + 1  # linha em branco entre as pastas
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def enviar(client, cabecalhos, conteudo, confirmar=False, ano=ANO):
    return client.post("/pastas/importar", headers=cabecalhos,
                       data={"ano": str(ano), "confirmar": str(confirmar).lower()},
                       files={"arquivo": ("pastas.xlsx", conteudo, XLSX)})


@pytest.fixture
def admin(headers):
    return headers("admin")


@pytest.fixture(autouse=True)
def preservar_ano_importado():
    """A importação substitui o ano; devolve os lançamentos do seed no fim."""
    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute("SELECT id_servo, id_pasta, momento, coordenador FROM servo_pasta WHERE ano = %s", (ANO,))
            antes = cursor.fetchall()
    yield
    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM servo_pasta WHERE ano = %s", (ANO,))
            for linha in antes:
                cursor.execute("INSERT INTO servo_pasta (id_servo, id_pasta, momento, coordenador, ano) "
                               "VALUES (%s, %s, %s, %s, %s)", (*linha, ANO))


def lancamentos(ano=ANO):
    with get_connection() as conn:
        with conn.cursor() as cursor:
            cursor.execute("""SELECT s.nome, p.nome, sp.momento, sp.coordenador FROM servo_pasta sp
                              JOIN servo s USING (id_servo) JOIN pasta p USING (id_pasta)
                              WHERE sp.ano = %s""", (ano,))
            return set(cursor.fetchall())


# ── Leitura da planilha ───────────────────────────────────────────────────────

def test_le_coordenador_membros_e_momentos():
    lidas = ler_planilha(planilha([
        {"funcao": "Música", "e": 1, "p": 1, "ec": ["Ana"], "em": ["Bia", "Caio"], "eo": "01 Coordenador + 02 Jovens",
         "pc": ["Ana"], "pm": ["Bia"]},
        {"funcao": "Cozinha", "e": 1, "em": ["Duda", "Casal Fulano", "Edu"], "vermelho": ["Edu"]},
        {"funcao": "Palestra", "e": 1, "em": ["Fábio", "Gabi"], "eo": "02 Jovens", "mesclar": {"e"}},
        {"funcao": "Coordenador Geral", "e": 1, "p": 1, "ec": ["Hugo", "Iara"], "eo": "02 Coordenadores",
         "mesclar": {"e"}, "pc": ["Hugo", "Iara"], "pm": ["Xxx e Xxx"], "amarelo": ["Xxx e Xxx"]},
    ]))
    musica, cozinha, palestra, geral = lidas

    assert musica.funcao == "Música"
    assert musica.equipes["encontro"].coordenadores == ["Ana"]
    assert musica.equipes["encontro"].membros == ["Bia", "Caio"]
    assert musica.equipes["pos"].membros == ["Bia"]
    assert musica.equipes["encontro"].marcada and musica.equipes["pos"].marcada

    # Apostolado (vermelho) e casais ficam de fora
    assert cozinha.equipes["encontro"].membros == ["Duda"]
    assert set(cozinha.equipes["encontro"].ignorados) == {"Casal Fulano", "Edu"}
    assert not cozinha.equipes["pos"].marcada

    # Nomes centralizados: "02 Jovens" → membros; "02 Coordenadores" → coordenadores
    assert palestra.equipes["encontro"].membros == ["Fábio", "Gabi"]
    assert palestra.equipes["encontro"].coordenadores == []
    assert geral.equipes["encontro"].coordenadores == ["Hugo", "Iara"]
    assert geral.equipes["pos"].ignorados == ["Xxx e Xxx"]


def test_arquivo_que_nao_e_planilha(client, admin):
    resposta = enviar(client, admin, b"isto nao e um xlsx")
    assert resposta.status_code == 422


def test_planilha_sem_cabecalho(client, admin):
    wb = Workbook()
    wb.active["A1"] = "Outra coisa"
    buffer = io.BytesIO()
    wb.save(buffer)
    resposta = enviar(client, admin, buffer.getvalue())
    assert resposta.status_code == 422
    assert "Função" in resposta.json()["detail"]


# ── Prévia e gravação ─────────────────────────────────────────────────────────

PLANILHA_OK = [
    {"funcao": "Música", "e": 1, "p": 1, "ec": ["Teste Servo 1"], "em": ["Teste Servo 3"],
     "pc": ["Teste Servo 1"], "pm": ["Teste Servo 3", "Fulano Que Não Existe"]},
    {"funcao": "Tutor da Comunidade de Cursistas: Kenossis", "e": 1, "p": 1,
     "pc": ["Teste Servo 5"], "po": "01 Coordenador", "mesclar": {"p"}},
    {"funcao": "Convêniência (Lojinha)", "e": 1, "em": ["Teste Servo 4"]},
]


def test_previa_nao_grava(client, admin):
    antes = lancamentos()
    resposta = enviar(client, admin, planilha(PLANILHA_OK))
    assert resposta.status_code == 200, resposta.text
    corpo = resposta.json()

    assert corpo["erros"] == []
    assert corpo["gravados"] == 0
    assert lancamentos() == antes

    pastas = {p["funcao_planilha"]: p for p in corpo["pastas"]}
    # Nome com letra trocada: acha, mas pede para conferir
    kenosis = pastas["Tutor da Comunidade de Cursistas: Kenossis"]
    assert kenosis["pasta"] == "Tutor da Comunidade de Cursistas: Kenosis" and kenosis["conferir"]
    # Pasta que não existe no sistema e pessoa que não é servo: só aviso
    assert pastas["Convêniência (Lojinha)"]["pasta"] is None
    assert any("Lojinha" in a for a in corpo["avisos"])
    assert any("Fulano Que Não Existe" in a for a in corpo["avisos"])


def test_confirmar_substitui_o_ano(client, admin):
    resposta = enviar(client, admin, planilha(PLANILHA_OK), confirmar=True)
    assert resposta.status_code == 200, resposta.text
    assert resposta.json()["gravados"] == 5

    assert lancamentos() == {
        ("Teste Servo 1", "Música", "encontro", True),
        ("Teste Servo 3", "Música", "encontro", False),
        ("Teste Servo 1", "Música", "pos", True),
        ("Teste Servo 3", "Música", "pos", False),
        ("Teste Servo 5", "Tutor da Comunidade de Cursistas: Kenosis", "pos", True),
    }


def test_regra_quebrada_impede_gravar(client, admin):
    # Teste Servo 2 fez Ornamentação no ano atual (seed): não pode repetir no seguinte
    antes = lancamentos()
    resposta = enviar(client, admin, planilha([
        {"funcao": "Ornamentação", "p": 1, "pc": ["Teste Servo 2"], "pm": ["Teste Servo 1"]},
    ]), confirmar=True)
    corpo = resposta.json()

    assert resposta.status_code == 200
    assert corpo["gravados"] == 0
    assert len(corpo["erros"]) == 1 and "não pode ser repetida" in corpo["erros"][0]
    assert "Nada foi gravado" in corpo["message"]
    assert lancamentos() == antes

    equipe = corpo["pastas"][0]["equipes"][0]
    assert equipe["coordenadores"][0]["erro"]
    assert equipe["membros"][0]["erro"] is None


def test_pasta_no_momento_errado(client, admin):
    resposta = enviar(client, admin, planilha([{"funcao": "Cozinha", "p": 1, "pm": ["Teste Servo 1"]}]))
    assert any("só existe no Encontro" in e for e in resposta.json()["erros"])


def test_so_ano_atual_ou_seguinte(client, admin):
    assert enviar(client, admin, planilha(PLANILHA_OK), ano=ANO - 2).status_code == 422


# ── Quem pode ─────────────────────────────────────────────────────────────────

def test_so_admin_e_coordenacao_geral(client, headers):
    conteudo = planilha(PLANILHA_OK)
    assert enviar(client, headers("servo"), conteudo).status_code == 403
    assert enviar(client, headers(None), conteudo).status_code == 401
    assert enviar(client, headers("geral.manual"), conteudo).status_code == 200
